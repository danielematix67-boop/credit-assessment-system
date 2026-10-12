from io import BytesIO

import openpyxl
import pytest

from src.services.excel_scenario_loader import (
    ExcelScenarioError,
    _convert_value,
    load_excel_scenarios,
)


def _workbook_bytes(
    *, include_customer_profile: bool = True,
    minimum_regulatory_risk_grade: str = "Past Due",
) -> bytes:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append([
        "scenario_id", "revenue", "ebitda", "revenue_growth",
        "bank_debt_current", "bank_debt_1y_ago", "bank_debt_2y_ago",
        "bank_debt_3y_ago", "bank_debt_4y_ago",
    ])
    credit.append(["TEST-001", 1000000, 100000, 0.05, 120, 100, 90, 80, 70])

    if include_customer_profile:
        profile = workbook.create_sheet("Customer Profile")
        profile.append(
            [
                "scenario_id",
                "company_name",
                "minimum_regulatory_risk_grade",
                "ews_score_class",
                "active_ewis",
            ]
        )
        profile.append(
            [
                "TEST-001",
                "Test Company",
                minimum_regulatory_risk_grade,
                "YELLOW",
                "Revenue deterioration;Payment delay",
            ]
        )

    behavioural = workbook.create_sheet("Behavioural Analysis")
    behavioural.append(["scenario_id", "average_utilization"])
    behavioural.append(["TEST-001", 0.75])

    debt = workbook.create_sheet("Debt Sustainability")
    debt.append(["scenario_id", "debt_service", "cash_flow_available_for_debt_service"])
    debt.append(["TEST-001", 200000, 300000])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_load_excel_scenario_maps_to_domain_models() -> None:
    scenarios = load_excel_scenarios(_workbook_bytes())

    position, profile, behavioural, debt = scenarios["TEST-001"]

    assert position.position_id == "TEST-001"
    assert position.revenue == 1000000.0
    assert position.bank_debt_current == 120.0
    assert position.bank_debt_4y_ago == 70.0
    assert profile is not None
    assert profile.minimum_regulatory_risk_grade == "Past Due"
    assert profile.active_ewis == ["Revenue deterioration", "Payment delay"]
    assert behavioural is not None
    assert behavioural.average_utilization == 0.75
    assert debt is not None
    assert debt.debt_service == 200000.0


def test_optional_domain_sheet_can_be_omitted() -> None:
    scenarios = load_excel_scenarios(
        _workbook_bytes(include_customer_profile=False)
    )

    assert scenarios["TEST-001"][1] is None


def test_invalid_regulatory_risk_grade_is_rejected() -> None:
    with pytest.raises(ExcelScenarioError, match="Allowed values"):
        load_excel_scenarios(
            _workbook_bytes(minimum_regulatory_risk_grade="Substandard")
        )


def test_required_credit_position_sheet_is_rejected_when_missing() -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "Customer Profile"
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="Financial Analysis"):
        load_excel_scenarios(buffer.getvalue())


def test_invalid_workbook_is_rejected() -> None:
    with pytest.raises(ExcelScenarioError, match="valid .xlsx"):
        load_excel_scenarios(b"not an excel workbook")


def test_sheet_without_scenario_id_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["revenue"])
    credit.append([100])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="scenario_id"):
        load_excel_scenarios(buffer.getvalue())


def test_empty_sheet_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="is empty"):
        load_excel_scenarios(buffer.getvalue())


def test_unsupported_column_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "unsupported"])
    credit.append(["TEST-001", 1])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="unsupported column"):
        load_excel_scenarios(buffer.getvalue())


def test_duplicate_credit_scenario_id_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append(["TEST-001", 100])
    credit.append(["TEST-001", 200])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="duplicate scenario_id"):
        load_excel_scenarios(buffer.getvalue())


def test_unknown_scenario_id_in_optional_sheet_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append(["TEST-001", 100])
    behavioural = workbook.create_sheet("Behavioural Analysis")
    behavioural.append(["scenario_id", "average_utilization"])
    behavioural.append(["TEST-999", 0.75])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="unknown scenario_id"):
        load_excel_scenarios(buffer.getvalue())


def test_multiple_rows_for_same_scenario_in_optional_sheet_are_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append(["TEST-001", 100])
    behavioural = workbook.create_sheet("Behavioural Analysis")
    behavioural.append(["scenario_id", "average_utilization"])
    behavioural.append(["TEST-001", 0.75])
    behavioural.append(["TEST-001", 0.80])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="multiple rows"):
        load_excel_scenarios(buffer.getvalue())


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("generational_transition", "maybe", "Invalid boolean"),
        ("relationship_years", "2.5", "Invalid integer"),
        ("revenue", "abc", "Invalid numeric"),
        ("ews_score_class", "PURPLE", "Invalid ews_score_class"),
    ],
)
def test_invalid_typed_values_are_rejected(
    field: str, value: str, message: str
) -> None:
    with pytest.raises(ExcelScenarioError, match=message):
        _convert_value(field, value)


def test_blank_scenario_id_is_rejected() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append([None, 100])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ExcelScenarioError, match="blank scenario_id"):
        load_excel_scenarios(buffer.getvalue())


def test_blank_risk_grade_change_is_normalized_to_none() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append(["TEST-001", 100])
    profile = workbook.create_sheet("Customer Profile")
    profile.append(["scenario_id", "risk_grade_change"])
    profile.append(["TEST-001", "   "])
    buffer = BytesIO()
    workbook.save(buffer)

    scenarios = load_excel_scenarios(buffer.getvalue())

    assert scenarios["TEST-001"][1] is not None
    assert scenarios["TEST-001"][1].risk_grade_change is None



def test_one_to_many_sheets_preserve_each_shareholder_event_and_facility() -> None:
    workbook = openpyxl.Workbook()
    credit = workbook.active
    credit.title = "Financial Analysis"
    credit.append(["scenario_id", "revenue"])
    credit.append(["TEST-001", 100])

    profile = workbook.create_sheet("Customer Profile")
    profile.append(["scenario_id", "company_name", "founding_year"])
    profile.append(["TEST-001", "Test Company", 1998])

    shareholders = workbook.create_sheet("Shareholders")
    shareholders.append(["scenario_id", "shareholder_name", "age", "role", "ownership_percentage"])
    shareholders.append(["TEST-001", "Alice Example", 52, "Chair", 60])
    shareholders.append(["TEST-001", "Bob Example", 48, "Director", 40])

    events = workbook.create_sheet("Relevant Events")
    events.append(["scenario_id", "event_type", "event_date", "description", "status"])
    events.append(["TEST-001", "Protest", "2026-01-15", "Supplier protest", "Open"])
    events.append(["TEST-001", "Restructuring", "2024-06-30", "Debt rescheduled", "Closed"])

    facilities = workbook.create_sheet("Credit Facilities")
    facilities.append([
        "scenario_id", "facility_id", "facility_type", "approved_amount",
        "used_amount", "residual_debt", "maturity_date", "review_date",
        "unpaid_installments", "forborne",
    ])
    facilities.append(["TEST-001", "F-001", "Term loan", 500000, 300000, 280000, "2028-12-31", 2, 3, True])
    facilities.append(["TEST-001", "F-002", "Revolving", 200000, 150000, None, None, "2027-03-31", 0, False])

    buffer = BytesIO()
    workbook.save(buffer)

    scenarios = load_excel_scenarios(buffer.getvalue())
    _, customer_profile, _, _ = scenarios["TEST-001"]

    assert customer_profile is not None
    assert customer_profile.founding_year == 1998
    assert len(customer_profile.shareholder_details) == 2
    assert len(customer_profile.relevant_event_details) == 2
    assert len(customer_profile.credit_facility_details) == 2
    assert "Alice Example" in customer_profile.shareholder_details[0]
    assert "60% of the shares" in customer_profile.shareholder_details[0]
    assert "F-001" in customer_profile.credit_facility_details[0]
    assert "approved limit €500,000" in customer_profile.credit_facility_details[0]
    assert "maturity 31 December 2028" in customer_profile.credit_facility_details[0]
    assert "review date" not in customer_profile.credit_facility_details[0]
    assert "3 unpaid instalment(s)" in customer_profile.credit_facility_details[0]
    assert "F-002" in customer_profile.credit_facility_details[1]
    assert "review date 31 March 2027" in customer_profile.credit_facility_details[1]
    assert "event_date=" not in customer_profile.relevant_event_details[0]
    assert "15 January 2026" in customer_profile.relevant_event_details[0]
    assert "1900" not in " ".join(customer_profile.credit_facility_details)
